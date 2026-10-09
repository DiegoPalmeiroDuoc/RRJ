"""HospitalOps web application. Server-rendered views + JSON dashboard endpoint."""
import io
import os
import re
import secrets
import hashlib
from datetime import timedelta
from decimal import Decimal, InvalidOperation
from contextlib import asynccontextmanager
from datetime import datetime
from typing import Any

from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.exceptions import RequestValidationError
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from app.database import get_db, init_db, SessionLocal
from app.models import (
    Hospital, Department, User, Asset, Ticket, TicketComment,
    InventoryItem, InventoryMovement, StockCount, StockCountLine,
    CRMContact, CRMOpportunity, AuditLog, ProviderItem, TechnicianStock, ProviderMovement, LoginAttempt, SessionEpoch, now,
)
from app.security import hash_password, verify_password

BASE = os.path.dirname(__file__)
PRODUCTION = os.environ.get('APP_ENV', 'development').lower() == 'production' or bool(os.environ.get('RAILWAY_ENVIRONMENT_NAME'))
if PRODUCTION:
    if len(os.environ.get('SESSION_SECRET', '')) < 40 or os.environ.get('SESSION_SECRET', '').startswith(('CAMBIA_', 'SOLO-')):
        raise RuntimeError('Producción requiere SESSION_SECRET aleatorio de 40+ caracteres')
    if os.environ.get('SESSION_HTTPS_ONLY', 'false').lower() != 'true':
        raise RuntimeError('En producción SESSION_HTTPS_ONLY=true es obligatorio')
    if not os.environ.get('DATABASE_URL', '').startswith(('postgresql://', 'postgresql+psycopg://', 'postgres://')):
        raise RuntimeError('En producción es obligatorio PostgreSQL')
    if not os.environ.get('ALLOWED_HOSTS', '').strip():
        raise RuntimeError('En producción define ALLOWED_HOSTS')


@asynccontextmanager
async def lifespan(_app):
    init_db()
    yield


app = FastAPI(title='HospitalOps', docs_url=None, redoc_url=None, lifespan=lifespan)
app.add_middleware(
    SessionMiddleware,
    secret_key=os.environ.get('SESSION_SECRET', 'SOLO-DESARROLLO-CAMBIA-ESTA-CLAVE'),
    session_cookie='hospitalops_session',
    same_site='lax',
    https_only=os.environ.get('SESSION_HTTPS_ONLY', 'false').lower() == 'true',
    max_age=60 * 60 * 10,
)
from starlette.middleware.trustedhost import TrustedHostMiddleware
if PRODUCTION:
    app.add_middleware(TrustedHostMiddleware, allowed_hosts=[h.strip() for h in os.environ['ALLOWED_HOSTS'].split(',')])

@app.middleware('http')
async def secure_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    response.headers['Permissions-Policy'] = 'camera=(self), microphone=(), geolocation=()'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; font-src 'self'; object-src 'none'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    if PRODUCTION:
        response.headers['Strict-Transport-Security'] = 'max-age=31536000; includeSubDomains'
    if request.url.path not in ('/static/styles.css', '/static/app.js'):
        response.headers['Cache-Control'] = 'no-store'
    return response

app.mount('/static', StaticFiles(directory=os.path.join(BASE, 'static')), name='static')
templates = Jinja2Templates(directory=os.path.join(BASE, 'templates'))

ROLES = {
    'superadmin': 'Administrador de plataforma',
    'coordinador_global': 'Coordinador empresa',
    'tecnico_global': 'Técnico empresa',
    'admin_cliente': 'Administrador hospital',
    'coordinador': 'Coordinador',
    'tecnico': 'Técnico',
    'solicitante': 'Solicitante',
}
ADMINS = {'superadmin', 'admin_cliente'}
MANAGERS = {'superadmin', 'coordinador_global', 'admin_cliente', 'coordinador'}
STAFF = MANAGERS | {'tecnico', 'tecnico_global'}
PROVIDER_MANAGERS = {'superadmin', 'coordinador_global'}
PRIORITIES = ['baja', 'media', 'alta', 'critica']
TICKET_STATUSES = ['pendiente', 'abierto', 'en_proceso', 'resuelto', 'cerrado', 'rechazado']
STAGES = ['nuevo', 'contactado', 'propuesta', 'negociacion', 'ganado', 'perdido']


def redirect(path='/dashboard'):
    return RedirectResponse(path, status_code=303)


XLSX_MEDIA_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def xlsx_response(filename: str, sheet_title: str, headers: list[str], rows, money_columns=()):
    """Excel download. Text is always stored as text, so values like '=HYPERLINK(...)' never run as formulas."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E79')
    for row in rows:
        values = []
        for value in row:
            if isinstance(value, Decimal):
                value = float(value)
            elif isinstance(value, bool):
                value = 'Sí' if value else 'No'
            values.append('' if value is None else value)
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for index in money_columns:
        for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            cell.number_format = '#,##0.00'
    for index, column in enumerate(sheet.columns, start=1):
        width = max(len(str(cell.value)) for cell in column if cell.value is not None)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 50)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    return StreamingResponse(iter([output.getvalue()]), media_type=XLSX_MEDIA_TYPE,
                             headers={'Content-Disposition': f'attachment; filename="{filename}"'})


def flash(request: Request, message: str, level: str = 'success'):
    request.session['flash'] = {'message': message, 'level': level}


def csrf(request: Request):
    if 'csrf' not in request.session:
        request.session['csrf'] = secrets.token_urlsafe(30)
    return request.session['csrf']


async def form_data(request: Request) -> Any:
    data = await request.form()
    if data.get('_csrf', '') != request.session.get('csrf') or not data.get('_csrf'):
        raise HTTPException(status_code=403, detail='Token CSRF inválido. Recarga la página.')
    return data


def clean(data: Any, key: str, maxlen: int = 250, default: str = '') -> str:
    return str(data.get(key, default) or '').strip()[:maxlen]


def positive_int(data, key, default=0) -> int:
    try:
        result = int(data.get(key, default))
    except (ValueError, TypeError):
        raise HTTPException(400, f'Número inválido: {key}')
    if result < 0 or result > 2_147_483_647:
        raise HTTPException(400, f'El valor {key} está fuera de rango')
    return result


def money(data, key, default='0') -> Decimal:
    try:
        value = Decimal(str(data.get(key, default))).quantize(Decimal('0.01'))
        if not value.is_finite() or value < 0 or value > 999_999_999_999:
            raise ValueError()
        return value
    except (InvalidOperation, ValueError):
        raise HTTPException(400, f'Monto inválido: {key}')


def chosen_id(data, key):
    raw = clean(data, key, 25)
    if not raw:
        return None
    try:
        number = int(raw)
        if number <= 0:
            raise ValueError()
        return number
    except ValueError:
        raise HTTPException(400, f'Identificador inválido: {key}')


def actor(request: Request, db: Session) -> User:
    user_id = request.session.get('uid')
    if not user_id:
        raise HTTPException(401, 'Debes iniciar sesión')
    user = db.get(User, user_id)
    if not user or not user.active:
        request.session.clear()
        raise HTTPException(401, 'Cuenta desactivada o sesión expirada')
    if user.approval != 'approved':
        raise HTTPException(403, 'Tu cuenta aún no está aprobada')
    security = db.get(SessionEpoch, user.id)
    if not security or request.session.get('epoch') != security.epoch:
        request.session.clear()
        raise HTTPException(401, 'Sesión revocada o expirada')
    if user.role in ('superadmin', 'tecnico_global', 'coordinador_global') and user.hospital_id is not None:
        raise HTTPException(403, 'El perfil de empresa no puede estar asignado a un hospital')
    if user.hospital_id and (not user.hospital or not user.hospital.active):
        request.session.clear()
        raise HTTPException(403, 'Hospital desactivado')
    return user


def hospital_for(request: Request, db: Session, user: User) -> Hospital:
    if user.role not in ('superadmin', 'coordinador_global', 'tecnico_global'):
        if not user.hospital_id:
            raise HTTPException(403, 'Usuario sin hospital')
        return user.hospital
    desired = request.session.get('hospital_id')
    hospital = db.get(Hospital, desired) if desired else None
    if not hospital or not hospital.active:
        hospital = db.scalar(select(Hospital).where(Hospital.active.is_(True)).order_by(Hospital.id))
    if not hospital:
        raise HTTPException(409, 'Primero registra un hospital desde Clientes')
    return hospital


def require(request: Request, db: Session, roles=None, module=None, allow_no_hospital=False):
    user = actor(request, db)
    if user.role == 'tecnico_global' and module != 'tickets' and not allow_no_hospital:
        raise HTTPException(403, 'Acceso limitado a tickets asignados')
    if roles and user.role not in roles:
        raise HTTPException(403, 'No tienes permisos para esta operación')
    hospital = None
    if not allow_no_hospital and not (user.role == 'tecnico_global' and module == 'tickets'):
        hospital = hospital_for(request, db, user)
        if module and not getattr(hospital, f'{module}_enabled'):
            raise HTTPException(403, 'Módulo desactivado en este hospital')
    return user, hospital


def scoped(db: Session, model, entity_id: int, hospital: Hospital):
    found = db.scalar(select(model).where(model.id == entity_id, model.hospital_id == hospital.id))
    if not found:
        raise HTTPException(404, 'Registro inexistente en este hospital')
    return found


def template(request: Request, db: Session, file: str, **context):
    user_id = request.session.get('uid')
    user = db.get(User, user_id) if user_id else None
    hospital = None
    if user and user.active and user.approval == 'approved' and user.role != 'tecnico_global':
        try:
            hospital = hospital_for(request, db, user)
        except HTTPException:
            pass
    hospitals = db.scalars(select(Hospital).order_by(Hospital.name)).all() if user and user.role in PROVIDER_MANAGERS else []
    payload = {
        'request': request, 'user': user, 'hospital': hospital,
        'hospitals': hospitals, 'csrf_token': csrf(request),
        'flash': request.session.pop('flash', None), 'roles': ROLES,
        'is_admin': bool(user and user.role in ADMINS),
        'is_provider': bool(user and user.role in PROVIDER_MANAGERS),
        'is_provider_technician': bool(user and user.role == 'tecnico_global'),
        'is_manager': bool(user and user.role in MANAGERS),
        'is_staff': bool(user and user.role in STAFF),
        'stages': STAGES, 'priorities': PRIORITIES, 'ticket_statuses': TICKET_STATUSES,
    }
    payload.update(context)
    return templates.TemplateResponse(request, file, payload)


def log(db: Session, user: User, hospital_id, action, entity, entity_id, detail=''):
    db.add(AuditLog(hospital_id=hospital_id, actor_id=user.id, action=action,
                    entity=entity, entity_id=entity_id, detail=detail[:250]))


def safe_flush(db: Session, message='El registro ya existe'):
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, message)


def unique_commit(db: Session, message='El registro ya existe'):
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, message)


def ensure_hospital_user(db, user_id, hospital, allowed_roles=None, required=False):
    if user_id is None:
        if required:
            raise HTTPException(400, 'Debes seleccionar un usuario')
        return None
    user = db.scalar(select(User).where(User.id == user_id, User.hospital_id == hospital.id,
                                        User.active.is_(True), User.approval == 'approved'))
    if not user or (allowed_roles and user.role not in allowed_roles):
        raise HTTPException(400, 'Usuario inválido para este hospital')
    return user


@app.exception_handler(HTTPException)
async def http_error(request: Request, exc: HTTPException):
    if exc.status_code == 401:
        return redirect('/login')
    with SessionLocal() as db:
        response = template(request, db, 'error.html', code=exc.status_code, detail=exc.detail)
        response.status_code = exc.status_code
        return response


@app.exception_handler(RequestValidationError)
async def validation_error(request: Request, exc: RequestValidationError):
    with SessionLocal() as db:
        response = template(request, db, 'error.html', code=422, detail='Datos de formulario inválidos')
        response.status_code = 422
        return response


@app.get('/', include_in_schema=False)
def home(request: Request):
    return redirect('/dashboard' if request.session.get('uid') else '/login')


@app.get('/health', include_in_schema=False)
def health():
    return {'status': 'ok'}


FAKE_HASH = hash_password('EstaContrasenaNoEsValida!2026')

@app.get('/ready', include_in_schema=False)
def ready(db: Session = Depends(get_db)):
    try:
        db.execute(text('SELECT 1'))
    except Exception:
        raise HTTPException(503, 'Base de datos no disponible')
    return {'status': 'ready'}


@app.get('/login', response_class=HTMLResponse)
def login_screen(request: Request, db: Session = Depends(get_db)):
    if request.session.get('uid'):
        account = db.get(User, request.session.get('uid'))
        security = db.get(SessionEpoch, account.id) if account else None
        if account and security and request.session.get('epoch') == security.epoch:
            return redirect('/dashboard' if account.approval == 'approved' else '/pending')
        request.session.clear()
    return template(request, db, 'login.html', registration_enabled=not PRODUCTION or os.getenv('ALLOW_SELF_REGISTRATION', 'false').lower() == 'true')


def login_identifier(request, email):
    ip = (request.client.host or 'unknown') if request.client else 'unknown'
    digest = hashlib.sha256(email.encode('utf-8')).hexdigest()
    return ip[:48], digest


def login_blocked(db, ip, digest):
    since = now() - timedelta(minutes=15)
    by_ip = db.scalar(select(func.count(LoginAttempt.id)).where(LoginAttempt.ip == ip, LoginAttempt.created_at >= since))
    by_email = db.scalar(select(func.count(LoginAttempt.id)).where(LoginAttempt.email_digest == digest, LoginAttempt.created_at >= since))
    return by_ip >= 200 or by_email >= 8


@app.post('/login')
async def login(request: Request, db: Session = Depends(get_db)):
    data = await form_data(request)
    email = clean(data, 'email', 160).lower()
    ip, digest = login_identifier(request, email)
    if login_blocked(db, ip, digest):
        raise HTTPException(429, 'Demasiados intentos. Intenta nuevamente en 15 minutos')
    user = db.scalar(select(User).where(User.email == email))
    valid = verify_password(clean(data, 'password', 500), user.password_hash) if user else verify_password(clean(data, 'password', 500), FAKE_HASH)
    if not user or not valid:
        db.add(LoginAttempt(ip=ip, email_digest=digest))
        db.commit()
        flash(request, 'Correo o contraseña incorrectos', 'error')
        return redirect('/login')
    if not user.active:
        flash(request, 'Cuenta desactivada; contacta al administrador', 'error')
        return redirect('/login')
    if user.hospital_id and (not user.hospital or not user.hospital.active):
        flash(request, 'Hospital desactivado', 'error')
        return redirect('/login')
    db.query(LoginAttempt).filter(LoginAttempt.email_digest == digest).delete(synchronize_session=False)
    db.commit()
    request.session.clear()  # regeneración lógica de sesión para no conservar datos anteriores
    security = db.get(SessionEpoch, user.id)
    if security is None:
        security = SessionEpoch(user_id=user.id, epoch=1)
        db.add(security)
        db.commit()
    request.session['uid'] = user.id
    request.session['epoch'] = security.epoch
    csrf(request)
    return redirect('/pending' if user.approval != 'approved' else '/dashboard')


@app.post('/logout')
async def logout(request: Request):
    await form_data(request)
    request.session.clear()
    return redirect('/login')


@app.get('/register', response_class=HTMLResponse)
def register_screen(request: Request, db: Session = Depends(get_db)):
    if PRODUCTION and os.getenv('ALLOW_SELF_REGISTRATION', 'false').lower() != 'true':
        raise HTTPException(404, 'Registro público deshabilitado')
    hospitals = db.scalars(select(Hospital).where(Hospital.active.is_(True)).order_by(Hospital.name)).all()
    return template(request, db, 'register.html', public_hospitals=hospitals)


@app.post('/register')
async def register(request: Request, db: Session = Depends(get_db)):
    if PRODUCTION and os.getenv('ALLOW_SELF_REGISTRATION', 'false').lower() != 'true':
        raise HTTPException(404, 'Registro público deshabilitado')
    data = await form_data(request)
    email = clean(data, 'email', 160).lower()
    name = clean(data, 'name', 150)
    password = clean(data, 'password', 500)
    hospital_id = chosen_id(data, 'hospital_id')
    hospital = db.scalar(select(Hospital).where(Hospital.id == hospital_id, Hospital.active.is_(True)))
    if not name or not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', email) or len(password) < 10 or not hospital:
        flash(request, 'Completa el nombre, correo, hospital y contraseña (mínimo 10 caracteres)', 'error')
        return redirect('/register')
    user = User(name=name, email=email, password_hash=hash_password(password),
                hospital_id=hospital.id, role='solicitante', approval='pending',
                job_title=clean(data, 'job_title', 120))
    db.add(user)
    unique_commit(db, 'El correo ya está registrado')
    request.session.clear()
    request.session['uid'] = user.id
    csrf(request)
    flash(request, 'Solicitud enviada; un administrador debe aprobar tu acceso')
    return redirect('/pending')


@app.get('/pending', response_class=HTMLResponse)
def pending(request: Request, db: Session = Depends(get_db)):
    user = db.get(User, request.session.get('uid')) if request.session.get('uid') else None
    if not user:
        return redirect('/login')
    if user.approval == 'approved':
        security = db.get(SessionEpoch, user.id)
        if not security or request.session.get('epoch') != security.epoch:
            request.session.clear()
            return redirect('/login')
        return redirect('/dashboard')
    return template(request, db, 'pending.html', pending_user=user)


@app.get('/perfil', response_class=HTMLResponse)
def profile(request: Request, db: Session = Depends(get_db)):
    require(request, db, allow_no_hospital=True)
    return template(request, db, 'profile.html')


@app.post('/perfil/password')
async def change_password(request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, allow_no_hospital=True)
    data = await form_data(request)
    old = clean(data, 'current_password', 500)
    new = clean(data, 'new_password', 500)
    if not verify_password(old, user.password_hash):
        raise HTTPException(400, 'Contraseña actual incorrecta')
    if len(new) < 12 or len(new) > 128 or new == old:
        raise HTTPException(400, 'La nueva contraseña debe tener entre 12 y 128 caracteres y ser distinta')
    user.password_hash = hash_password(new)
    security = db.get(SessionEpoch, user.id)
    if security is not None:
        security.epoch += 1
    log(db, user, user.hospital_id, 'actualizar_contrasena', 'usuario', user.id)
    db.commit()
    request.session.clear()
    flash(request, 'Contraseña actualizada. Ingresa nuevamente')
    return redirect('/login')


@app.post('/switch-hospital')
async def switch_hospital(request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    data = await form_data(request)
    hospital_id = chosen_id(data, 'hospital_id')
    hospital = db.scalar(select(Hospital).where(Hospital.id == hospital_id, Hospital.active.is_(True)))
    if not hospital:
        raise HTTPException(400, 'Hospital no disponible')
    request.session['hospital_id'] = hospital.id
    return redirect('/dashboard')


@app.get('/dashboard', response_class=HTMLResponse)
def dashboard(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, allow_no_hospital=True)
    if user.role == 'tecnico_global':
        return redirect('/mis-tickets')
    hospital = hospital_for(request, db, user)
    if user.role == 'tecnico_global':
        return redirect('/mis-tickets')
    query = select(Ticket).where(Ticket.hospital_id == hospital.id)
    if user.role == 'solicitante':
        query = query.where(Ticket.requester_id == user.id)
    if user.role == 'tecnico':
        query = query.where(or_(Ticket.assignee_id == user.id, Ticket.requester_id == user.id))
    tickets = db.scalars(query.order_by(Ticket.created_at.desc()).limit(6)).all()
    events = db.scalars(select(AuditLog).where(AuditLog.hospital_id == hospital.id)
                        .order_by(AuditLog.created_at.desc()).limit(7)).all() if user.role in ADMINS else []
    return template(request, db, 'dashboard.html', tickets=tickets, events=events)


@app.get('/api/dashboard')
def api_dashboard(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, allow_no_hospital=True)
    if user.role == 'tecnico_global':
        raise HTTPException(403)
    hospital = hospital_for(request, db, user)
    ticket_where = [Ticket.hospital_id == hospital.id]
    if user.role == 'solicitante':
        ticket_where.append(Ticket.requester_id == user.id)
    if user.role == 'tecnico':
        ticket_where.append(or_(Ticket.assignee_id == user.id, Ticket.requester_id == user.id))
    tickets = db.scalars(select(Ticket).where(*ticket_where)).all()
    response = {
        'tickets_total': len(tickets),
        'tickets_active': sum(t.status in ('pendiente', 'abierto', 'en_proceso') for t in tickets),
        'tickets_pending': sum(t.status == 'pendiente' for t in tickets),
        'tickets_resolved': sum(t.status in ('resuelto', 'cerrado') for t in tickets),
    }
    if user.role in STAFF:
        items = db.scalars(select(InventoryItem).where(InventoryItem.hospital_id == hospital.id,
                                                      InventoryItem.active.is_(True))).all()
        response['stock_low'] = sum(i.stock <= i.min_stock for i in items)
        response['stock_value'] = round(sum(float(i.gross_cost) * i.stock for i in items))
    return JSONResponse(response, headers={'Cache-Control': 'no-store'})


# --- HOSPITALES (administrador de plataforma) ---
@app.get('/hospitals', response_class=HTMLResponse)
def hospitals_list(request: Request, db: Session = Depends(get_db)):
    require(request, db, {'superadmin'}, allow_no_hospital=True)
    items = db.scalars(select(Hospital).order_by(Hospital.name)).all()
    return template(request, db, 'hospitals.html', items=items)


@app.get('/hospitals/new', response_class=HTMLResponse)
def hospital_new(request: Request, db: Session = Depends(get_db)):
    require(request, db, {'superadmin'}, allow_no_hospital=True)
    return template(request, db, 'hospital_form.html', item=None)


@app.get('/hospitals/{id}/edit', response_class=HTMLResponse)
def hospital_edit(id: int, request: Request, db: Session = Depends(get_db)):
    require(request, db, {'superadmin'}, allow_no_hospital=True)
    item = db.get(Hospital, id)
    if not item:
        raise HTTPException(404, 'Hospital inexistente')
    return template(request, db, 'hospital_form.html', item=item)


def update_hospital(item, data):
    item.name = clean(data, 'name', 160)
    item.code = clean(data, 'code', 30).upper()
    item.tax_id = clean(data, 'tax_id', 40)
    item.address = clean(data, 'address', 250)
    item.email = clean(data, 'email', 160)
    item.phone = clean(data, 'phone', 50)
    if not item.name or not item.code:
        raise HTTPException(400, 'Nombre y código son obligatorios')


@app.post('/hospitals/new')
async def hospital_create(request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, {'superadmin'}, allow_no_hospital=True)
    data = await form_data(request)
    item = Hospital()
    update_hospital(item, data)
    db.add(item)
    safe_flush(db, 'Hospital duplicado')
    log(db, user, item.id, 'crear', 'hospital', item.id, item.name)
    unique_commit(db, 'Ya existe un hospital con ese nombre o código')
    flash(request, 'Hospital creado')
    return redirect('/hospitals')


@app.post('/hospitals/{id}/edit')
async def hospital_update(id: int, request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, {'superadmin'}, allow_no_hospital=True)
    item = db.get(Hospital, id)
    if not item:
        raise HTTPException(404)
    data = await form_data(request)
    update_hospital(item, data)
    log(db, user, id, 'editar', 'hospital', id, item.name)
    unique_commit(db, 'Nombre o código duplicado')
    flash(request, 'Hospital actualizado')
    return redirect('/hospitals')


@app.post('/hospitals/{id}/toggle')
async def hospital_toggle(id: int, request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, {'superadmin'}, allow_no_hospital=True)
    await form_data(request)
    item = db.get(Hospital, id)
    if not item:
        raise HTTPException(404)
    item.active = not item.active
    log(db, user, id, 'activar' if item.active else 'desactivar', 'hospital', id, item.name)
    db.commit()
    flash(request, 'Estado del hospital actualizado')
    return redirect('/hospitals')


# --- USUARIOS / APROBACIÓN ---
@app.get('/users', response_class=HTMLResponse)
def users_list(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    users = db.scalars(select(User).where(User.hospital_id == hospital.id).order_by(User.name)).all()
    return template(request, db, 'users.html', items=users)


@app.get('/users/new', response_class=HTMLResponse)
def users_new(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    deps = db.scalars(select(Department).where(Department.hospital_id == hospital.id)).all()
    return template(request, db, 'user_form.html', item=None, deps=deps)


@app.get('/users/{id}/edit', response_class=HTMLResponse)
def user_edit(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    item = scoped(db, User, id, hospital)
    deps = db.scalars(select(Department).where(Department.hospital_id == hospital.id)).all()
    return template(request, db, 'user_form.html', item=item, deps=deps)


def update_user(item, data, hospital, db, current, is_new=False):
    role = clean(data, 'role', 25)
    if role not in {'admin_cliente', 'coordinador', 'tecnico', 'solicitante'}:
        raise HTTPException(400, 'Rol no permitido')
    if item.id == current.id and role != current.role:
        raise HTTPException(403, 'No puedes cambiar tu propio rol')
    item.name = clean(data, 'name', 150)
    item.email = clean(data, 'email', 160).lower()
    item.job_title = clean(data, 'job_title', 120)
    item.role = role
    item.hospital_id = hospital.id
    dep_id = chosen_id(data, 'department_id')
    if dep_id and not db.scalar(select(Department).where(Department.id == dep_id,
                                                        Department.hospital_id == hospital.id)):
        raise HTTPException(400, 'Área inválida')
    item.department_id = dep_id
    password = clean(data, 'password', 500)
    if is_new and len(password) < (12 if PRODUCTION else 10):
        raise HTTPException(400, 'La contraseña debe tener al menos 12 caracteres en producción')
    if password:
        if len(password) < (12 if PRODUCTION else 10):
            raise HTTPException(400, 'La contraseña debe tener al menos 12 caracteres en producción')
        item.password_hash = hash_password(password)
        if not is_new:
            security = db.get(SessionEpoch, item.id)
            if security:
                security.epoch += 1
    if not item.name or not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', item.email):
        raise HTTPException(400, 'Nombre y correo válido son obligatorios')
    item.approval = 'approved'


@app.post('/users/new')
async def users_create(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    data = await form_data(request)
    item = User()
    update_user(item, data, hospital, db, user, True)
    db.add(item)
    safe_flush(db, 'Correo ya registrado')
    log(db, user, hospital.id, 'crear', 'usuario', item.id, item.name)
    unique_commit(db, 'El correo ya está registrado')
    flash(request, 'Usuario creado')
    return redirect('/users')


@app.post('/users/{id}/edit')
async def users_update(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    item = scoped(db, User, id, hospital)
    data = await form_data(request)
    update_user(item, data, hospital, db, user)
    log(db, user, hospital.id, 'editar', 'usuario', item.id, item.name)
    unique_commit(db, 'El correo ya está registrado')
    flash(request, 'Usuario actualizado')
    return redirect('/users')


@app.post('/users/{id}/toggle')
async def users_toggle(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    await form_data(request)
    item = scoped(db, User, id, hospital)
    if item.id == user.id:
        raise HTTPException(403, 'No puedes desactivar tu propia cuenta')
    item.active = not item.active
    log(db, user, hospital.id, 'activar' if item.active else 'desactivar', 'usuario', id, item.name)
    db.commit()
    flash(request, 'Estado de usuario actualizado')
    return redirect('/users')


@app.get('/approvals', response_class=HTMLResponse)
def approvals(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    items = db.scalars(select(User).where(User.hospital_id == hospital.id,
                                           User.approval == 'pending').order_by(User.created_at)).all()
    return template(request, db, 'approvals.html', items=items)


@app.post('/approvals/{id}/{action}')
async def approval_decision(id: int, action: str, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    await form_data(request)
    if action not in ('approve', 'reject'):
        raise HTTPException(400)
    item = scoped(db, User, id, hospital)
    if item.approval != 'pending':
        raise HTTPException(409, 'Solicitud ya procesada')
    item.approval = 'approved' if action == 'approve' else 'rejected'
    log(db, user, hospital.id, action, 'usuario', id, item.email)
    db.commit()
    flash(request, 'Solicitud procesada')
    return redirect('/approvals')


@app.get('/departments', response_class=HTMLResponse)
def departments(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    items = db.scalars(select(Department).where(Department.hospital_id == hospital.id).order_by(Department.name)).all()
    return template(request, db, 'departments.html', items=items)


@app.post('/departments')
async def departments_create(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    data = await form_data(request)
    name = clean(data, 'name', 120)
    if not name:
        raise HTTPException(400, 'El área no puede estar vacía')
    item = Department(hospital_id=hospital.id, name=name)
    db.add(item)
    safe_flush(db, 'Área duplicada')
    log(db, user, hospital.id, 'crear', 'area', item.id, item.name)
    unique_commit(db, 'Ya existe un área con ese nombre')
    flash(request, 'Área creada')
    return redirect('/departments')


# --- EQUIPOS / ACTIVOS ---
@app.get('/assets', response_class=HTMLResponse)
def assets(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, STAFF, 'inventory')
    items = db.scalars(select(Asset).where(Asset.hospital_id == hospital.id).order_by(Asset.name)).all()
    return template(request, db, 'assets.html', items=items)


@app.get('/assets/new', response_class=HTMLResponse)
def assets_new(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    return template(request, db, 'asset_form.html', item=None)


@app.get('/assets/{id}/edit', response_class=HTMLResponse)
def assets_edit(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    return template(request, db, 'asset_form.html', item=scoped(db, Asset, id, hospital))


@app.post('/assets/{id}/save')
@app.post('/assets/new')
async def assets_save(request: Request, id: int | None = None, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    data = await form_data(request)
    item = scoped(db, Asset, id, hospital) if id is not None else Asset(hospital_id=hospital.id)
    item.code = clean(data, 'code', 70).upper()
    item.name = clean(data, 'name', 150)
    item.category = clean(data, 'category', 90)
    item.serial_number = clean(data, 'serial_number', 100)
    item.location = clean(data, 'location', 150)
    item.active = data.get('active') == 'on'
    if not item.code or not item.name:
        raise HTTPException(400, 'Código y nombre obligatorios')
    db.add(item)
    safe_flush(db, 'Código de activo duplicado')
    log(db, user, hospital.id, 'editar' if id else 'crear', 'activo', item.id, item.code)
    unique_commit(db, 'El código del activo ya existe')
    flash(request, 'Equipo registrado')
    return redirect('/assets')

# --- VISTA MULTIHOSPITAL PARA PERSONAL DE EMPRESA ---
@app.get('/operaciones/tickets', response_class=HTMLResponse)
def provider_tickets(request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    hospital_id = request.query_params.get('hospital_id', '')
    status = request.query_params.get('status', '')
    query = select(Ticket).join(Hospital, Hospital.id == Ticket.hospital_id)
    if hospital_id.isdigit():
        query = query.where(Ticket.hospital_id == int(hospital_id))
    if status in TICKET_STATUSES:
        query = query.where(Ticket.status == status)
    page = max(1, min(int(request.query_params.get('page', '1')) if request.query_params.get('page','1').isdigit() else 1, 100000))
    count = db.scalar(select(func.count()).select_from(query.subquery()))
    items = db.scalars(query.order_by(Ticket.created_at.desc()).offset((page-1)*50).limit(50)).all()
    hospitals = db.scalars(select(Hospital).order_by(Hospital.name)).all()
    return template(request, db, 'provider_tickets.html', items=items, clients=hospitals,
                    status=status, chosen_hospital=hospital_id, page=page, total=count)


@app.post('/operaciones/abrir/{id}')
async def provider_open_ticket(id: int, request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    await form_data(request)
    item = db.get(Ticket, id)
    if not item:
        raise HTTPException(404)
    request.session['hospital_id'] = item.hospital_id
    return redirect(f'/tickets/{id}')


@app.get('/mis-tickets', response_class=HTMLResponse)
def my_global_tickets(request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, {'tecnico_global'}, allow_no_hospital=True)
    items = db.scalars(select(Ticket).join(Hospital, Hospital.id == Ticket.hospital_id).where(
        Ticket.assignee_id == user.id, Hospital.active.is_(True))
                       .order_by(Ticket.created_at.desc()).limit(150)).all()
    return template(request, db, 'my_global_tickets.html', items=items)


# --- TICKETS ---
def ticket_for_user(db, id, hospital, user):
    if user.role == 'tecnico_global':
        item = db.get(Ticket, id)
        if not item or item.assignee_id != user.id:
            raise HTTPException(404, 'Ticket no asignado')
        active = db.scalar(select(Hospital.id).where(Hospital.id == item.hospital_id,
                                                    Hospital.active.is_(True), Hospital.tickets_enabled.is_(True)))
        if not active:
            raise HTTPException(404, 'Hospital no disponible')
        return item
    return scoped(db, Ticket, id, hospital)


def ticket_visible(user, ticket):
    if user.role == 'solicitante' and ticket.requester_id != user.id:
        raise HTTPException(404, 'Ticket no encontrado')
    if user.role == 'tecnico_global' and ticket.assignee_id != user.id:
        raise HTTPException(404, 'Ticket no asignado')
    if user.role == 'tecnico' and ticket.assignee_id != user.id and ticket.requester_id != user.id:
        raise HTTPException(404, 'Ticket no asignado a tu cuenta')


@app.get('/tickets', response_class=HTMLResponse)
def tickets(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, module='tickets')
    if user.role == 'tecnico_global':
        return redirect('/mis-tickets')
    query = select(Ticket).where(Ticket.hospital_id == hospital.id)
    if user.role == 'solicitante':
        query = query.where(Ticket.requester_id == user.id)
    if user.role == 'tecnico':
        query = query.where(or_(Ticket.assignee_id == user.id, Ticket.requester_id == user.id))
    status = request.query_params.get('status', '')
    q = request.query_params.get('q', '').strip()[:100]
    if status in TICKET_STATUSES:
        query = query.where(Ticket.status == status)
    if q:
        query = query.where(or_(Ticket.title.ilike(f'%{q}%'), Ticket.description.ilike(f'%{q}%')))
    items = db.scalars(query.order_by(Ticket.created_at.desc()).limit(250)).all()
    return template(request, db, 'tickets.html', items=items, status=status, q=q)


@app.get('/tickets/new', response_class=HTMLResponse)
def ticket_new(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, module='tickets')
    if user.role == 'tecnico_global':
        raise HTTPException(403)
    assets = db.scalars(select(Asset).where(Asset.hospital_id == hospital.id, Asset.active.is_(True))).all()
    return template(request, db, 'ticket_form.html', assets=assets)


@app.post('/tickets/new')
async def ticket_create(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, module='tickets')
    if user.role == 'tecnico_global':
        raise HTTPException(403)
    data = await form_data(request)
    title, description = clean(data, 'title', 200), clean(data, 'description', 5000)
    if not title or not description:
        raise HTTPException(400, 'Título y descripción son obligatorios')
    asset_id = chosen_id(data, 'asset_id')
    if asset_id:
        scoped(db, Asset, asset_id, hospital)
    priority = clean(data, 'priority', 20)
    if priority not in PRIORITIES:
        raise HTTPException(400, 'Prioridad inválida')
    item = Ticket(hospital_id=hospital.id, requester_id=user.id, title=title,
                  description=description, category=clean(data, 'category', 90) or 'Soporte',
                  priority=priority, location=clean(data, 'location', 150), asset_id=asset_id,
                  status='pendiente')
    db.add(item)
    db.flush()
    log(db, user, hospital.id, 'crear', 'ticket', item.id, title)
    db.commit()
    flash(request, f'Ticket #{item.id} enviado. Pendiente de confirmación administrativa')
    return redirect(f'/tickets/{item.id}')


@app.get('/tickets/{id}', response_class=HTMLResponse)
def ticket_detail(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, module='tickets')
    item = ticket_for_user(db, id, hospital, user)
    ticket_visible(user, item)
    query = select(TicketComment).where(TicketComment.ticket_id == id)
    if user.role == 'solicitante':
        query = query.where(TicketComment.internal.is_(False))
    comments = db.scalars(query.order_by(TicketComment.created_at)).all()
    assignees = db.scalars(select(User).where(or_(
                                               (User.hospital_id == hospital.id) & User.role.in_(['coordinador', 'tecnico']),
                                               (User.role == 'tecnico_global') & User.hospital_id.is_(None)),
                                               User.active.is_(True), User.approval == 'approved').order_by(User.name)).all() if user.role in MANAGERS else []
    return template(request, db, 'ticket_detail.html', item=item, comments=comments, assignees=assignees)


@app.post('/tickets/{id}/confirm')
async def ticket_confirm(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS, 'tickets')
    await form_data(request)
    item = ticket_for_user(db, id, hospital, user)
    if item.status != 'pendiente':
        raise HTTPException(409, 'El ticket ya fue confirmado o rechazado')
    item.status = 'abierto'
    item.confirmed_by_id = user.id
    item.confirmed_at = now()
    log(db, user, hospital.id, 'confirmar', 'ticket', id, item.title)
    db.commit()
    flash(request, 'Ticket confirmado, listo para asignar')
    return redirect(f'/tickets/{id}')


@app.post('/tickets/{id}/reject')
async def ticket_reject(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS, 'tickets')
    await form_data(request)
    item = ticket_for_user(db, id, hospital, user)
    if item.status != 'pendiente':
        raise HTTPException(409, 'Solo se rechazan tickets pendientes')
    item.status = 'rechazado'
    log(db, user, hospital.id, 'rechazar', 'ticket', id, item.title)
    db.commit()
    flash(request, 'Ticket rechazado')
    return redirect(f'/tickets/{id}')


@app.post('/tickets/{id}/assign')
async def ticket_assign(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'tickets')
    data = await form_data(request)
    item = ticket_for_user(db, id, hospital, user)
    if item.status in ('pendiente', 'rechazado', 'cerrado'):
        raise HTTPException(409, 'Confirma el ticket antes de asignar')
    user_id = chosen_id(data, 'assignee_id')
    if user_id:
        candidate = db.get(User, user_id)
        if not candidate or not candidate.active or candidate.approval != 'approved' or not (
            (candidate.hospital_id == hospital.id and candidate.role in ('tecnico', 'coordinador')) or
            (candidate.role == 'tecnico_global' and candidate.hospital_id is None)
        ):
            raise HTTPException(400, 'Técnico no autorizado')
    item.assignee_id = user_id
    log(db, user, hospital.id, 'asignar', 'ticket', id, f'Responsable: {user_id}')
    db.commit()
    flash(request, 'Responsable actualizado')
    return redirect(f'/tickets/{id}')


@app.post('/tickets/{id}/status')
async def ticket_update_status(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, STAFF, 'tickets')
    data = await form_data(request)
    item = ticket_for_user(db, id, hospital, user)
    ticket_visible(user, item)
    new_status = clean(data, 'status', 25)
    allowed = {
        'abierto': {'en_proceso', 'resuelto'},
        'en_proceso': {'abierto', 'resuelto'},
        'resuelto': {'en_proceso', 'cerrado'},
        'cerrado': {'abierto'} if user.role in ADMINS else set(),
    }
    if new_status not in allowed.get(item.status, set()):
        raise HTTPException(400, 'Transición de estado no permitida')
    if user.role in ('tecnico', 'tecnico_global') and (item.assignee_id != user.id or new_status not in ('en_proceso', 'resuelto', 'abierto')):
        raise HTTPException(403)
    item.status = new_status
    log(db, user, item.hospital_id, 'estado', 'ticket', id, new_status)
    db.commit()
    flash(request, 'Estado actualizado')
    return redirect(f'/tickets/{id}')


@app.post('/tickets/{id}/close')
async def ticket_requester_close(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, module='tickets')
    await form_data(request)
    item = ticket_for_user(db, id, hospital, user)
    if user.role != 'solicitante' or item.requester_id != user.id or item.status != 'resuelto':
        raise HTTPException(403, 'Solo el solicitante puede aceptar la resolución')
    item.status = 'cerrado'
    log(db, user, hospital.id, 'aceptar_resolucion', 'ticket', id, item.title)
    db.commit()
    flash(request, 'Resolución aceptada. Ticket cerrado')
    return redirect(f'/tickets/{id}')


@app.post('/tickets/{id}/comments')
async def ticket_comment(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, module='tickets')
    data = await form_data(request)
    item = ticket_for_user(db, id, hospital, user)
    ticket_visible(user, item)
    body = clean(data, 'body', 4000)
    if not body:
        raise HTTPException(400, 'El comentario está vacío')
    internal = data.get('internal') == 'on'
    if internal and user.role not in STAFF:
        raise HTTPException(403)
    db.add(TicketComment(ticket_id=id, author_id=user.id, body=body, internal=internal))
    log(db, user, item.hospital_id, 'comentar', 'ticket', id, 'Nota interna' if internal else 'Comentario')
    db.commit()
    flash(request, 'Comentario publicado')
    return redirect(f'/tickets/{id}#comments')


# --- INVENTARIO / KARDEX ---
@app.get('/inventory', response_class=HTMLResponse)
def inventory(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, STAFF, 'inventory')
    query = select(InventoryItem).where(InventoryItem.hospital_id == hospital.id)
    q = request.query_params.get('q', '').strip()[:100]
    if q:
        query = query.where(or_(InventoryItem.name.ilike(f'%{q}%'), InventoryItem.sku.ilike(f'%{q}%'),
                                InventoryItem.barcode.ilike(f'%{q}%')))
    items = db.scalars(query.order_by(InventoryItem.name).limit(500)).all()
    return template(request, db, 'inventory.html', items=items, q=q)


@app.get('/inventory/new', response_class=HTMLResponse)
def inventory_new(request: Request, db: Session = Depends(get_db)):
    require(request, db, MANAGERS, 'inventory')
    return template(request, db, 'inventory_form.html', item=None)


@app.get('/inventory/{id}/edit', response_class=HTMLResponse)
def inventory_edit(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    return template(request, db, 'inventory_form.html', item=scoped(db, InventoryItem, id, hospital))


@app.post('/inventory/new')
@app.post('/inventory/{id}/edit')
async def inventory_save(request: Request, id: int | None = None, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    data = await form_data(request)
    item = scoped(db, InventoryItem, id, hospital) if id is not None else InventoryItem(hospital_id=hospital.id)
    item.name = clean(data, 'name', 160)
    item.sku = clean(data, 'sku', 60).upper()
    item.barcode = clean(data, 'barcode', 100)
    item.category = clean(data, 'category', 100)
    item.location = clean(data, 'location', 150)
    item.gross_cost = money(data, 'gross_cost')
    item.sale_price = money(data, 'sale_price')
    item.min_stock = positive_int(data, 'min_stock')
    item.active = data.get('active') == 'on'
    if not item.name or not item.sku:
        raise HTTPException(400, 'SKU y nombre son obligatorios')
    if id is None:
        item.stock = positive_int(data, 'stock')
        db.add(item)
        safe_flush(db, 'SKU ya registrado')
        if item.stock:
            db.add(InventoryMovement(hospital_id=hospital.id, item_id=item.id, actor_id=user.id,
                                     kind='inicial', delta=item.stock, resulting_stock=item.stock,
                                     note='Carga inicial'))
    log(db, user, hospital.id, 'editar' if id else 'crear', 'producto', item.id, item.sku)
    unique_commit(db, 'El SKU ya existe en este hospital')
    flash(request, 'Producto guardado. Para modificar stock utiliza Kardex / Movimientos')
    return redirect('/inventory')


@app.get('/inventory/{id}', response_class=HTMLResponse)
def inventory_detail(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, STAFF, 'inventory')
    item = scoped(db, InventoryItem, id, hospital)
    moves = db.scalars(select(InventoryMovement).where(InventoryMovement.hospital_id == hospital.id,
                                                        InventoryMovement.item_id == id)
                       .order_by(InventoryMovement.created_at.desc()).limit(100)).all()
    return template(request, db, 'inventory_detail.html', item=item, moves=moves)


@app.post('/inventory/{id}/move')
async def inventory_move(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    data = await form_data(request)
    item = db.scalar(select(InventoryItem).where(InventoryItem.id == id, InventoryItem.hospital_id == hospital.id).with_for_update())
    if not item:
        raise HTTPException(404, 'Producto no encontrado en este hospital')
    if not item.active:
        raise HTTPException(409, 'No se mueve stock de artículos inactivos')
    kind = clean(data, 'kind', 25)
    amount = positive_int(data, 'amount')
    if kind not in ('entrada', 'salida', 'ajuste') or (kind != 'ajuste' and amount == 0):
        raise HTTPException(400, 'Movimiento inválido')
    delta = amount if kind == 'entrada' else -amount if kind == 'salida' else amount - item.stock
    if item.stock + delta < 0:
        raise HTTPException(409, 'Stock insuficiente')
    ticket_id = chosen_id(data, 'ticket_id')
    if ticket_id:
        scoped(db, Ticket, ticket_id, hospital)
    note = clean(data, 'note', 250)
    if not note:
        raise HTTPException(400, 'Debes indicar el motivo del movimiento')
    item.stock += delta
    db.add(InventoryMovement(hospital_id=hospital.id, item_id=id, actor_id=user.id,
                             kind=kind, delta=delta, resulting_stock=item.stock, note=note, ticket_id=ticket_id))
    log(db, user, hospital.id, kind, 'producto', id, f'{delta:+d} unidades — {note}')
    db.commit()
    flash(request, f'Stock actualizado: {item.stock} unidades')
    return redirect(f'/inventory/{id}')


@app.get('/inventory/export/xlsx')
def inventory_export(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    items = db.scalars(select(InventoryItem).where(InventoryItem.hospital_id == hospital.id)
                       .order_by(InventoryItem.sku)).all()
    return xlsx_response('inventario.xlsx', 'Inventario',
        ['SKU', 'Código barras', 'Producto', 'Categoría', 'Ubicación',
         'Precio bruto', 'Precio venta', 'Stock', 'Stock mínimo', 'Activo'],
        [(i.sku, i.barcode, i.name, i.category, i.location, i.gross_cost, i.sale_price,
          i.stock, i.min_stock, i.active) for i in items],
        money_columns=(6, 7))


@app.get('/counts', response_class=HTMLResponse)
def counts(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, STAFF, 'inventory')
    items = db.scalars(select(StockCount).where(StockCount.hospital_id == hospital.id)
                       .order_by(StockCount.created_at.desc())).all()
    return template(request, db, 'counts.html', items=items)


@app.post('/counts/new')
async def count_create(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    data = await form_data(request)
    item = StockCount(hospital_id=hospital.id, created_by_id=user.id, note=clean(data, 'note', 200))
    db.add(item)
    safe_flush(db)
    inventory_items = db.scalars(select(InventoryItem).where(InventoryItem.hospital_id == hospital.id,
                                                             InventoryItem.active.is_(True)).order_by(InventoryItem.name)).all()
    if not inventory_items:
        raise HTTPException(400, 'Primero registra productos activos')
    for inv in inventory_items:
        db.add(StockCountLine(stock_count_id=item.id, item_id=inv.id, expected_stock=inv.stock))
    log(db, user, hospital.id, 'crear', 'conteo', item.id, item.note)
    db.commit()
    flash(request, 'Conteo abierto; registra el stock físico de cada producto')
    return redirect(f'/counts/{item.id}')


@app.get('/counts/{id}', response_class=HTMLResponse)
def count_detail(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, STAFF, 'inventory')
    item = scoped(db, StockCount, id, hospital)
    lines = db.scalars(select(StockCountLine).where(StockCountLine.stock_count_id == id)
                       .order_by(StockCountLine.id)).all()
    return template(request, db, 'count_detail.html', item=item, lines=lines)


@app.post('/counts/{id}/close')
async def count_close(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'inventory')
    data = await form_data(request)
    item = scoped(db, StockCount, id, hospital)
    if item.status != 'abierto':
        raise HTTPException(409, 'El conteo ya está cerrado')
    lines = db.scalars(select(StockCountLine).where(StockCountLine.stock_count_id == id)).all()
    changes = []
    for line in lines:
        current = db.scalar(select(InventoryItem).where(InventoryItem.id == line.item_id, InventoryItem.hospital_id == hospital.id).with_for_update())
        if not current:
            raise HTTPException(404, 'Producto del conteo inexistente')
        if current.stock != line.expected_stock:
            raise HTTPException(409, f'Stock de {current.sku} cambió desde la apertura; crea un nuevo conteo')
        actual = positive_int(data, f'actual_{line.id}', line.expected_stock)
        changes.append((line, current, actual))
    for line, current, actual in changes:
        line.actual_stock = actual
        delta = actual - current.stock
        if delta:
            current.stock = actual
            db.add(InventoryMovement(hospital_id=hospital.id, item_id=current.id, actor_id=user.id,
                                     kind='conteo', delta=delta, resulting_stock=actual,
                                     note=f'Ajuste inventario físico #{id}'))
    item.status = 'cerrado'
    item.closed_at = now()
    log(db, user, hospital.id, 'cerrar', 'conteo', id, f'{len(lines)} líneas verificadas')
    db.commit()
    flash(request, 'Inventario físico conciliado y cerrado')
    return redirect(f'/counts/{id}')


# --- CRM ---
@app.get('/crm', response_class=HTMLResponse)
def crm(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'crm')
    opportunities = db.scalars(select(CRMOpportunity).where(CRMOpportunity.hospital_id == hospital.id)
                               .order_by(CRMOpportunity.created_at.desc())).all()
    by_stage = {stage: [o for o in opportunities if o.stage == stage] for stage in STAGES}
    total_value = sum(float(o.value) for o in opportunities if o.stage not in ('perdido',))
    return template(request, db, 'crm.html', by_stage=by_stage, total_value=total_value, count=len(opportunities))


@app.get('/crm/contacts', response_class=HTMLResponse)
def crm_contacts(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'crm')
    items = db.scalars(select(CRMContact).where(CRMContact.hospital_id == hospital.id)
                       .order_by(CRMContact.name)).all()
    return template(request, db, 'crm_contacts.html', items=items)


@app.get('/crm/contacts/new', response_class=HTMLResponse)
def crm_contact_new(request: Request, db: Session = Depends(get_db)):
    require(request, db, MANAGERS, 'crm')
    return template(request, db, 'crm_contact_form.html', item=None)


@app.get('/crm/contacts/{id}/edit', response_class=HTMLResponse)
def crm_contact_edit(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'crm')
    return template(request, db, 'crm_contact_form.html', item=scoped(db, CRMContact, id, hospital))


@app.post('/crm/contacts/new')
@app.post('/crm/contacts/{id}/edit')
async def crm_contact_save(request: Request, id: int | None = None, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'crm')
    data = await form_data(request)
    item = scoped(db, CRMContact, id, hospital) if id is not None else CRMContact(hospital_id=hospital.id)
    item.name = clean(data, 'name', 150)
    item.email = clean(data, 'email', 160)
    item.phone = clean(data, 'phone', 50)
    item.job_title = clean(data, 'job_title', 150)
    item.department = clean(data, 'department', 130)
    item.notes = clean(data, 'notes', 3000)
    if not item.name:
        raise HTTPException(400, 'Nombre requerido')
    db.add(item)
    safe_flush(db)
    log(db, user, hospital.id, 'editar' if id else 'crear', 'contacto', item.id, item.name)
    db.commit()
    flash(request, 'Contacto guardado')
    return redirect('/crm/contacts')


@app.get('/crm/opportunities/new', response_class=HTMLResponse)
def opportunity_new(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'crm')
    contacts = db.scalars(select(CRMContact).where(CRMContact.hospital_id == hospital.id)).all()
    owners = db.scalars(select(User).where(User.hospital_id == hospital.id, User.role.in_(list(MANAGERS - {'superadmin'})),
                                           User.active.is_(True), User.approval == 'approved')).all()
    return template(request, db, 'crm_opportunity_form.html', item=None, contacts=contacts, owners=owners)


@app.get('/crm/opportunities/{id}/edit', response_class=HTMLResponse)
def opportunity_edit(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'crm')
    item = scoped(db, CRMOpportunity, id, hospital)
    contacts = db.scalars(select(CRMContact).where(CRMContact.hospital_id == hospital.id)).all()
    owners = db.scalars(select(User).where(User.hospital_id == hospital.id, User.role.in_(list(MANAGERS - {'superadmin'})),
                                           User.active.is_(True), User.approval == 'approved')).all()
    return template(request, db, 'crm_opportunity_form.html', item=item, contacts=contacts, owners=owners)


@app.post('/crm/opportunities/new')
@app.post('/crm/opportunities/{id}/edit')
async def opportunity_save(request: Request, id: int | None = None, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'crm')
    data = await form_data(request)
    item = scoped(db, CRMOpportunity, id, hospital) if id is not None else CRMOpportunity(hospital_id=hospital.id)
    item.name = clean(data, 'name', 190)
    item.value = money(data, 'value')
    item.stage = clean(data, 'stage', 30)
    item.expected_close = clean(data, 'expected_close', 12)
    item.notes = clean(data, 'notes', 4000)
    if item.stage not in STAGES or not item.name:
        raise HTTPException(400, 'Nombre o etapa inválidos')
    contact_id = chosen_id(data, 'contact_id')
    if contact_id:
        scoped(db, CRMContact, contact_id, hospital)
    item.contact_id = contact_id
    owner_id = chosen_id(data, 'owner_id')
    if owner_id:
        ensure_hospital_user(db, owner_id, hospital, {'admin_cliente', 'coordinador'})
    item.owner_id = owner_id
    db.add(item)
    safe_flush(db)
    log(db, user, hospital.id, 'editar' if id else 'crear', 'oportunidad', item.id, item.name)
    db.commit()
    flash(request, 'Oportunidad guardada')
    return redirect('/crm')


@app.post('/crm/opportunities/{id}/stage')
async def opportunity_stage(id: int, request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, MANAGERS, 'crm')
    data = await form_data(request)
    item = scoped(db, CRMOpportunity, id, hospital)
    stage = clean(data, 'stage', 30)
    if stage not in STAGES:
        raise HTTPException(400, 'Etapa inválida')
    item.stage = stage
    log(db, user, hospital.id, 'etapa', 'oportunidad', id, stage)
    db.commit()
    flash(request, 'Oportunidad movida')
    return redirect('/crm')


# --- PERSONALIZACIÓN / AUDITORÍA ---
@app.get('/settings', response_class=HTMLResponse)
def settings(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    return template(request, db, 'settings.html')


@app.post('/settings')
async def settings_update(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    data = await form_data(request)
    color = clean(data, 'color', 7)
    if not re.fullmatch(r'#[A-Fa-f0-9]{6}', color):
        raise HTTPException(400, 'Color hexadecimal inválido')
    portal_name = clean(data, 'portal_name', 90)
    if not portal_name:
        raise HTTPException(400, 'Nombre del portal requerido')
    hospital.color = color
    hospital.portal_name = portal_name
    hospital.tickets_enabled = data.get('tickets_enabled') == 'on'
    hospital.inventory_enabled = data.get('inventory_enabled') == 'on'
    hospital.crm_enabled = data.get('crm_enabled') == 'on'
    log(db, user, hospital.id, 'configurar', 'hospital', hospital.id, 'Imagen corporativa y módulos')
    db.commit()
    flash(request, 'Personalización actualizada')
    return redirect('/settings')


@app.get('/audit', response_class=HTMLResponse)
def audit(request: Request, db: Session = Depends(get_db)):
    user, hospital = require(request, db, ADMINS)
    items = db.scalars(select(AuditLog).where(AuditLog.hospital_id == hospital.id)
                       .order_by(AuditLog.created_at.desc()).limit(300)).all()
    return template(request, db, 'audit.html', items=items)


# Database initialization is handled by the FastAPI lifespan context above.


# --- EQUIPO PROVEEDOR: usuarios sin pertenencia a hospitales ---
@app.get('/equipo', response_class=HTMLResponse)
def provider_team(request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    items = db.scalars(select(User).where(User.role.in_(['coordinador_global', 'tecnico_global']))
                       .order_by(User.name)).all()
    return template(request, db, 'provider_team.html', items=items)


@app.post('/equipo/crear')
async def provider_add_user(request: Request, db: Session = Depends(get_db)):
    actor_user, _ = require(request, db, {'superadmin'}, allow_no_hospital=True)
    data = await form_data(request)
    name, email = clean(data, 'name', 150), clean(data, 'email', 160).lower()
    role, password = clean(data, 'role', 25), clean(data, 'password', 500)
    if (not name or not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', email)
        or role not in ('coordinador_global', 'tecnico_global') or len(password) < 12):
        raise HTTPException(400, 'Nombre, correo, rol y contraseña de 12+ caracteres son obligatorios')
    item = User(name=name, email=email, role=role, password_hash=hash_password(password),
                hospital_id=None, approval='approved', active=True)
    db.add(item)
    safe_flush(db, 'Correo ya registrado')
    log(db, actor_user, None, 'crear', 'equipo_empresa', item.id, name)
    unique_commit(db)
    flash(request, 'Integrante creado. Debe cambiar su contraseña en Mi cuenta')
    return redirect('/equipo')


@app.post('/equipo/{id}/reset-password')
async def provider_reset_password(id: int, request: Request, db: Session = Depends(get_db)):
    manager, _ = require(request, db, {'superadmin'}, allow_no_hospital=True)
    data = await form_data(request)
    item = db.get(User, id)
    new_password = clean(data, 'password', 500)
    if (not item or item.role not in ('tecnico_global', 'coordinador_global')
        or item.hospital_id is not None):
        raise HTTPException(404)
    if len(new_password) < 12 or len(new_password) > 128:
        raise HTTPException(400, 'Contraseña inválida: 12 a 128 caracteres')
    item.password_hash = hash_password(new_password)
    security = db.get(SessionEpoch, item.id)
    if security:
        security.epoch += 1
    log(db, manager, None, 'reset_contrasena', 'equipo_empresa', id)
    db.commit()
    flash(request, 'Contraseña restablecida; se revocaron las sesiones previas')
    return redirect('/equipo')


@app.post('/equipo/{id}/activar')
async def provider_toggle_user(id: int, request: Request, db: Session = Depends(get_db)):
    manager, _ = require(request, db, {'superadmin'}, allow_no_hospital=True)
    await form_data(request)
    user = db.get(User, id)
    if not user or user.role not in ('tecnico_global', 'coordinador_global') or user.hospital_id is not None:
        raise HTTPException(404, 'Usuario no encontrado')
    user.active = not user.active
    log(db, manager, None, 'cambiar_acceso', 'equipo_empresa', id, 'Activo' if user.active else 'Inactivo')
    db.commit()
    flash(request, 'Acceso actualizado')
    return redirect('/equipo')


# --- BODEGA CENTRAL DE LA EMPRESA PRESTADORA ---
@app.get('/bodega', response_class=HTMLResponse)
def provider_warehouse(request: Request, db: Session = Depends(get_db)):
    manager, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    items = db.scalars(select(ProviderItem).order_by(ProviderItem.name).limit(500)).all()
    users = db.scalars(select(User).where(User.role == 'tecnico_global', User.active.is_(True),
                                          User.approval == 'approved').order_by(User.name)).all()
    allocations = db.scalars(select(TechnicianStock).where(TechnicianStock.quantity > 0).order_by(TechnicianStock.id.desc()).limit(500)).all()
    movements = db.scalars(select(ProviderMovement).order_by(ProviderMovement.id.desc()).limit(75)).all()
    return template(request, db, 'provider_warehouse.html', items=items, users=users,
                    allocations=allocations, movements=movements)


@app.post('/bodega/crear')
async def provider_item_create(request: Request, db: Session = Depends(get_db)):
    manager, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    data = await form_data(request)
    sku, name = clean(data, 'sku', 60).upper(), clean(data, 'name', 160)
    if not sku or not name:
        raise HTTPException(400, 'SKU y nombre son obligatorios')
    item = ProviderItem(sku=sku, name=name, category=clean(data, 'category', 100) or 'General',
                        location=clean(data, 'location', 120) or 'Bodega central',
                        gross_cost=money(data, 'gross_cost'), sale_price=money(data, 'sale_price'),
                        stock=positive_int(data, 'stock'), min_stock=positive_int(data, 'min_stock'))
    db.add(item)
    safe_flush(db, 'SKU ya existe')
    if item.stock:
        db.add(ProviderMovement(item_id=item.id, actor_id=manager.id, kind='inicial',
                                quantity=item.stock, note='Carga inicial'))
    log(db, manager, None, 'crear', 'bodega_producto', item.id, sku)
    unique_commit(db)
    flash(request, 'Producto de bodega registrado')
    return redirect('/bodega')


@app.post('/bodega/{id}/editar')
async def provider_item_update(id: int, request: Request, db: Session = Depends(get_db)):
    manager, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    data = await form_data(request)
    item = db.get(ProviderItem, id)
    if not item:
        raise HTTPException(404)
    item.name = clean(data, 'name', 160)
    item.sku = clean(data, 'sku', 60).upper()
    item.category = clean(data, 'category', 100) or 'General'
    item.location = clean(data, 'location', 120) or 'Bodega central'
    item.gross_cost, item.sale_price = money(data, 'gross_cost'), money(data, 'sale_price')
    item.min_stock = positive_int(data, 'min_stock')
    item.active = data.get('active') == 'on'
    if not item.name or not item.sku:
        raise HTTPException(400, 'SKU y nombre obligatorios')
    log(db, manager, None, 'editar', 'bodega_producto', item.id, item.name)
    unique_commit(db, 'El SKU está siendo utilizado')
    return redirect('/bodega')


def lock_provider_item(db, id, require_active=True):
    item = db.scalar(select(ProviderItem).where(ProviderItem.id == id).with_for_update())
    if not item or (require_active and not item.active):
        raise HTTPException(404, 'Producto inexistente o inactivo')
    return item


def valid_provider_technician(db, technician_id):
    tech = db.get(User, technician_id) if technician_id else None
    if not tech or tech.role != 'tecnico_global' or tech.hospital_id is not None or not tech.active or tech.approval != 'approved':
        raise HTTPException(400, 'Técnico inválido')
    return tech


@app.post('/bodega/{id}/movimiento')
async def provider_warehouse_movement(id: int, request: Request, db: Session = Depends(get_db)):
    manager, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    data = await form_data(request)
    item = lock_provider_item(db, id)
    kind = clean(data, 'kind', 20)
    quantity = positive_int(data, 'quantity')
    note = clean(data, 'note', 250)
    if not note or kind not in ('entrada', 'salida', 'ajuste'):
        raise HTTPException(400, 'Movimiento inválido')
    if quantity == 0 and kind != 'ajuste':
        raise HTTPException(400, 'Cantidad inválida')
    delta = quantity if kind == 'entrada' else -quantity if kind == 'salida' else quantity - item.stock
    if item.stock + delta < 0:
        raise HTTPException(409, 'Stock insuficiente')
    item.stock += delta
    db.add(ProviderMovement(item_id=id, actor_id=manager.id, kind=kind,
                            quantity=delta, note=note))
    log(db, manager, None, kind, 'bodega_producto', id, f'{delta:+d} — {note}')
    db.commit()
    return redirect('/bodega')


@app.post('/bodega/{id}/asignar')
async def provider_assign_stock(id: int, request: Request, db: Session = Depends(get_db)):
    manager, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    data = await form_data(request)
    item = lock_provider_item(db, id)
    tech_id = chosen_id(data, 'technician_id')
    valid_provider_technician(db, tech_id)
    qty = positive_int(data, 'quantity')
    if qty <= 0 or qty > item.stock:
        raise HTTPException(409, 'Cantidad superior al stock disponible')
    # Se requiere bloqueo del stock central; la asignación es única por producto y técnico.
    allocation = db.scalar(select(TechnicianStock).where(TechnicianStock.item_id == id,
        TechnicianStock.technician_id == tech_id).with_for_update())
    if allocation is None:
        allocation = TechnicianStock(item_id=id, technician_id=tech_id, quantity=0)
        db.add(allocation)
    allocation.quantity += qty
    item.stock -= qty
    db.add(ProviderMovement(item_id=id, technician_id=tech_id, actor_id=manager.id,
                            kind='asignacion', quantity=qty, note='Entrega a técnico'))
    log(db, manager, None, 'entrega', 'bodega_producto', id, f'{qty} unidades -> técnico {tech_id}')
    unique_commit(db, 'Asignación concurrente: vuelve a intentar')
    return redirect('/bodega')


@app.post('/bodega/{id}/devolver')
async def provider_return_stock(id: int, request: Request, db: Session = Depends(get_db)):
    manager, _ = require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    data = await form_data(request)
    item = lock_provider_item(db, id, require_active=False)
    tech_id = chosen_id(data, 'technician_id')
    allocation = db.scalar(select(TechnicianStock).where(TechnicianStock.item_id == id,
            TechnicianStock.technician_id == tech_id).with_for_update())
    qty = positive_int(data, 'quantity')
    if not allocation or qty < 1 or qty > allocation.quantity:
        raise HTTPException(409, 'Técnico no tiene esa cantidad')
    allocation.quantity -= qty
    item.stock += qty
    db.add(ProviderMovement(item_id=id, technician_id=tech_id, actor_id=manager.id,
                            kind='devolucion', quantity=qty, note='Devolución a bodega'))
    log(db, manager, None, 'devolucion', 'bodega_producto', id, str(qty))
    db.commit()
    return redirect('/bodega')


@app.get('/mis-repuestos', response_class=HTMLResponse)
def provider_my_stock(request: Request, db: Session = Depends(get_db)):
    user, _ = require(request, db, {'tecnico_global'}, allow_no_hospital=True)
    items = db.scalars(select(TechnicianStock).where(TechnicianStock.technician_id == user.id,
                                                     TechnicianStock.quantity > 0)).all()
    assigned = db.scalars(select(Ticket).join(Hospital, Hospital.id == Ticket.hospital_id).where(
        Ticket.assignee_id == user.id, Hospital.active.is_(True), Ticket.status.in_(['abierto', 'en_proceso'])).order_by(Ticket.id.desc())).all()
    return template(request, db, 'provider_my_stock.html', items=items, tickets=assigned)


@app.post('/mis-repuestos/{id}/consumir')
async def provider_consume_stock(id: int, request: Request, db: Session = Depends(get_db)):
    technician, _ = require(request, db, {'tecnico_global'}, allow_no_hospital=True)
    data = await form_data(request)
    ticket_id = chosen_id(data, 'ticket_id')
    qty = positive_int(data, 'quantity')
    note = clean(data, 'note', 250)
    ticket = db.get(Ticket, ticket_id) if ticket_id else None
    if (not ticket or ticket.assignee_id != technician.id or ticket.status not in ('abierto', 'en_proceso')
        or not db.scalar(select(Hospital.id).where(Hospital.id == ticket.hospital_id, Hospital.active.is_(True)))):
        raise HTTPException(403, 'Solo puedes registrar consumo en tus tickets abiertos')
    if qty < 1 or not note:
        raise HTTPException(400, 'Cantidad y motivo obligatorios')
    item = lock_provider_item(db, id)
    allocation = db.scalar(select(TechnicianStock).where(TechnicianStock.item_id == id,
            TechnicianStock.technician_id == technician.id).with_for_update())
    if not allocation or allocation.quantity < qty:
        raise HTTPException(409, 'Stock de técnico insuficiente')
    allocation.quantity -= qty
    db.add(ProviderMovement(item_id=id, actor_id=technician.id, technician_id=technician.id,
                            ticket_id=ticket.id, hospital_id=ticket.hospital_id,
                            kind='consumo', quantity=qty, note=note))
    log(db, technician, ticket.hospital_id, 'consumo', 'bodega_producto', id, f'{qty} unidades en ticket {ticket.id}')
    db.commit()
    flash(request, 'Consumo registrado en ticket')
    return redirect('/mis-repuestos')


@app.get('/bodega/export/xlsx')
def provider_warehouse_export(request: Request, db: Session = Depends(get_db)):
    require(request, db, PROVIDER_MANAGERS, allow_no_hospital=True)
    items = db.scalars(select(ProviderItem).order_by(ProviderItem.sku)).all()
    return xlsx_response('bodega_empresa.xlsx', 'Bodega',
        ['SKU','Producto','Categoría','Ubicación','Costo bruto','Precio venta','Stock bodega','Stock mínimo'],
        [(item.sku,item.name,item.category,item.location,item.gross_cost,item.sale_price,
          item.stock,item.min_stock) for item in items],
        money_columns=(5, 6))
