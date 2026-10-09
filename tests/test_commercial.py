"""Acceptance/security regression tests for the service-provider edition."""
import io
from openpyxl import load_workbook
from sqlalchemy import select
from app.models import User, Ticket, ProviderItem, TechnicianStock, ProviderMovement, Hospital
from test_app import sandbox, signin, token, form_post


def sign_in_new(client, email, password):
    page = client.get('/login')
    return client.post('/login', data={'_csrf': token(page.text), 'email': email, 'password': password})


def add_technician(client, name='Pedro Operaciones', email='pedro@servicios.example'):
    return form_post(client, '/equipo/crear', {'name': name, 'email': email,
        'role':'tecnico_global', 'password':'PasswordUnicaComercial2026!'})


def test_superadmin_global_team_and_hospital_admin_cannot_escalate(sandbox):
    client, Session = sandbox
    signin(client, 'admin@norte.example')
    assert client.get('/equipo').status_code == 403
    assert form_post(client, '/users/new', {'name':'Invasor', 'email':'invasor@norte.example',
        'role':'coordinador_global', 'password':'ContrasenaExtensa2026!'}).status_code == 400
    with Session() as db:
        assert db.scalar(select(User).where(User.email == 'invasor@norte.example')) is None
    client.cookies.clear()
    signin(client, 'root@hospitalops.example')
    assert client.get('/equipo').status_code == 200
    assert add_technician(client).status_code == 303
    assert 'Pedro Operaciones' in client.get('/equipo').text
    with Session() as db:
        tech = db.scalar(select(User).where(User.email == 'pedro@servicios.example'))
        assert tech.role == 'tecnico_global' and tech.hospital_id is None


def test_multihospital_technician_sees_only_assigned(sandbox):
    client, Session = sandbox
    signin(client, 'root@hospitalops.example')
    add_technician(client)
    with Session() as db:
        tech = db.scalar(select(User).where(User.email == 'pedro@servicios.example'))
        tech_id = tech.id
        # Seed tickets 2 and 4 belong to different hospitals; the confirmed statuses are known.
        a, b = db.get(Ticket, 2), db.get(Ticket, 4)
        assert a.hospital_id != b.hospital_id
        a.assignee_id = b.assignee_id = tech_id
        db.commit()
    client.cookies.clear()
    sign_in_new(client, 'pedro@servicios.example', 'PasswordUnicaComercial2026!')
    assert client.get('/mis-tickets').status_code == 200
    assert client.get('/tickets/2').status_code == 200
    assert client.get('/tickets/4').status_code == 200
    assert client.get('/tickets/1').status_code == 404
    assert client.get('/inventory').status_code == 403
    assert client.get('/crm').status_code == 403
    assert client.get('/users').status_code == 403
    assert client.get('/operaciones/tickets').status_code == 403
    assert client.get('/equipo').status_code == 403
    assert client.get('/bodega').status_code == 403
    assert client.get('/mis-repuestos').status_code == 200
    assert client.get('/dashboard').status_code == 303


def test_provider_coordinator_uses_all_hospitals_without_hospital_admin_privileges(sandbox):
    client, Session = sandbox
    signin(client, 'root@hospitalops.example')
    assert form_post(client, '/equipo/crear', {'name':'Coord Central','email':'coord@empresa.example',
        'role':'coordinador_global','password':'PasswordUnicaComercial2026!'}).status_code == 303
    client.cookies.clear()
    sign_in_new(client, 'coord@empresa.example','PasswordUnicaComercial2026!')
    assert client.get('/operaciones/tickets').status_code == 200
    assert 'Hospital Demo Norte' in client.get('/operaciones/tickets').text
    assert 'Hospital Demo Sur' in client.get('/operaciones/tickets').text
    assert client.get('/bodega').status_code == 200
    assert client.get('/hospitals').status_code == 403
    assert client.get('/users').status_code == 403
    assert client.get('/settings').status_code == 403
    assert client.get('/equipo').status_code == 200
    assert form_post(client, '/equipo/crear', {'name':'No Autorizado','email':'no@empresa.example',
        'role':'tecnico_global','password':'PasswordUnicaComercial2026!'}).status_code == 403


def test_warehouse_assignment_consumption_and_audit(sandbox):
    client, Session = sandbox
    signin(client, 'root@hospitalops.example')
    add_technician(client)
    result = form_post(client, '/bodega/crear', {'sku':'REP-001','name':'Cable médico',
        'stock':'10', 'min_stock':'2','gross_cost':'2500','sale_price':'4200','location':'Central'})
    assert result.status_code == 303
    with Session() as db:
        item = db.scalar(select(ProviderItem).where(ProviderItem.sku == 'REP-001'))
        tech = db.scalar(select(User).where(User.email == 'pedro@servicios.example'))
        ticket = db.get(Ticket, 2)
        ticket.assignee_id = tech.id
        ticket.status = 'abierto'
        db.commit()
        item_id, tech_id = item.id, tech.id
    assert form_post(client, f'/bodega/{item_id}/asignar', {'technician_id': str(tech_id),
        'quantity':'4'}).status_code == 303
    with Session() as db:
        assert db.get(ProviderItem,item_id).stock == 6
        assert db.scalar(select(TechnicianStock).where(TechnicianStock.item_id==item_id)).quantity == 4
    assert form_post(client, f'/bodega/{item_id}/asignar', {'technician_id': str(tech_id),
        'quantity':'7'}).status_code == 409
    # A hospital requester must never view or manage the provider's stock.
    client.cookies.clear()
    signin(client, 'solicitante@norte.example')
    assert client.get('/bodega').status_code == 403
    assert client.get('/mis-repuestos').status_code == 403
    assert form_post(client, f'/bodega/{item_id}/movimiento', {'kind':'entrada','quantity':'99',
        'note':'Attack'}).status_code == 403
    client.cookies.clear()
    sign_in_new(client, 'pedro@servicios.example','PasswordUnicaComercial2026!')
    assert 'Cable médico' in client.get('/mis-repuestos').text
    assert form_post_no_dash(client, f'/mis-repuestos/{item_id}/consumir', {'ticket_id':'2',
        'quantity':'2','note':'Cambio de cable'}).status_code == 303
    assert form_post_no_dash(client, f'/mis-repuestos/{item_id}/consumir', {'ticket_id':'1',
        'quantity':'1','note':'No asignado'}).status_code == 403
    assert form_post_no_dash(client, f'/mis-repuestos/{item_id}/consumir', {'ticket_id':'2',
        'quantity':'3','note':'Más stock'}).status_code == 409
    with Session() as db:
        assert db.get(ProviderItem,item_id).stock == 6
        assert db.scalar(select(TechnicianStock).where(TechnicianStock.item_id==item_id)).quantity == 2
        movement = db.scalar(select(ProviderMovement).where(ProviderMovement.kind == 'consumo'))
        assert movement.ticket_id == 2 and movement.hospital_id == 1 and movement.quantity == 2
    client.cookies.clear()
    signin(client, 'root@hospitalops.example')
    assert form_post(client, f'/bodega/{item_id}/devolver', {'technician_id': str(tech_id),
        'quantity':'2'}).status_code == 303
    with Session() as db:
        assert db.get(ProviderItem, item_id).stock == 8
        assert db.scalar(select(TechnicianStock).where(TechnicianStock.item_id==item_id)).quantity == 0


def test_login_throttling_and_security_headers(sandbox):
    client, Session = sandbox
    for _ in range(8):
        page = client.get('/login')
        invalid = client.post('/login', data={'_csrf': token(page.text),
            'email':'admin@norte.example', 'password':'fallo'})
        assert invalid.status_code == 303
    page = client.get('/login')
    assert client.post('/login', data={'_csrf':token(page.text),
        'email':'admin@norte.example', 'password':'DemoSeguro2026!'}).status_code == 429
    assert 'Content-Security-Policy' in page.headers
    assert 'X-Frame-Options' in page.headers
    assert page.headers['Cache-Control'] == 'no-store'


def test_admin_password_change_invalidates_session(sandbox):
    client, Session = sandbox
    signin(client, 'root@hospitalops.example')
    result = form_post(client, '/perfil/password', {
        'current_password':'DemoSeguro2026!', 'new_password':'ContrasenaNuevaYExtensa2026!'})
    assert result.status_code == 303
    assert result.headers['location'] == '/login'
    assert client.get('/equipo').status_code == 303
    login = sign_in_new(client, 'root@hospitalops.example', 'ContrasenaNuevaYExtensa2026!')
    assert login.status_code == 303 and login.headers['location'] == '/dashboard'


def test_global_tech_cannot_self_create_ticket(sandbox):
    client, _ = sandbox
    signin(client, 'root@hospitalops.example')
    add_technician(client)
    client.cookies.clear()
    sign_in_new(client, 'pedro@servicios.example','PasswordUnicaComercial2026!')
    assert client.get('/tickets/new').status_code == 403
    assert form_post_no_dash(client, '/tickets/new', {'title':'A','description':'B','priority':'media'}).status_code == 403


def form_post_no_dash(client, route, data):
    page = client.get('/mis-tickets')
    return client.post(route, data={'_csrf':token(page.text), **data})


def test_password_change_revokes_existing_browser_sessions(sandbox):
    from fastapi.testclient import TestClient
    from app.main import app
    client, _ = sandbox
    signin(client, 'root@hospitalops.example')
    with TestClient(app, follow_redirects=False) as other:
        signin(other, 'root@hospitalops.example')
        assert other.get('/equipo').status_code == 200
        assert form_post(client, '/perfil/password', {
            'current_password':'DemoSeguro2026!', 'new_password':'PasswordRevocadaExtensa2026!'}) .status_code == 303
        assert other.get('/equipo').status_code == 303


def test_warehouse_xlsx_export_and_extra_isolation(sandbox):
    client, _ = sandbox
    signin(client, 'root@hospitalops.example')
    assert form_post(client, '/bodega/crear', {'sku':'SNEAK-01','name':'=HYPERLINK("x")',
        'stock':'1','gross_cost':'11','sale_price':'20'}).status_code == 303
    response = client.get('/bodega/export/xlsx')
    assert response.status_code == 200
    rows = {row[0].value: row for row in load_workbook(io.BytesIO(response.content)).active.iter_rows(min_row=2)}
    name_cell = rows['SNEAK-01'][1]
    # El texto se guarda como texto literal, nunca como fórmula ejecutable
    assert name_cell.value == '=HYPERLINK("x")' and name_cell.data_type == 's'
    client.cookies.clear()
    signin(client, 'admin@norte.example')
    assert client.get('/bodega/export/xlsx').status_code == 403
    assert client.get('/operaciones/tickets').status_code == 403


def test_signed_in_provider_tech_cannot_be_scoped_through_hospital_selector(sandbox):
    client, _ = sandbox
    signin(client, 'root@hospitalops.example')
    add_technician(client)
    client.cookies.clear()
    sign_in_new(client, 'pedro@servicios.example', 'PasswordUnicaComercial2026!')
    assert form_post_no_dash(client, '/switch-hospital', {'hospital_id':'2'}).status_code == 403
    assert form_post_no_dash(client, '/equipo/crear', {
        'name':'Hacker','email':'hacker@example.org','password':'ContraseñaExtensa2026!',
        'role':'coordinador_global'}).status_code == 403


def test_real_assignment_of_global_technician_via_ticket_ui(sandbox):
    client, Session = sandbox
    signin(client, 'root@hospitalops.example')
    assert add_technician(client).status_code == 303
    with Session() as db:
        tech = db.scalar(select(User).where(User.email == 'pedro@servicios.example'))
        global_id = tech.id
    client.cookies.clear()
    signin(client, 'admin@norte.example')
    assert form_post(client, '/tickets/2/confirm', {}).status_code == 303
    html = client.get('/tickets/2').text
    assert 'Pedro Operaciones' in html
    assert form_post(client, '/tickets/2/assign', {'assignee_id':str(global_id)}).status_code == 303
    with Session() as db:
        assert db.get(Ticket, 2).assignee_id == global_id
    client.cookies.clear()
    sign_in_new(client, 'pedro@servicios.example','PasswordUnicaComercial2026!')
    assert client.get('/tickets/2').status_code == 200
    assert form_post_no_dash(client, '/tickets/2/status', {'status':'en_proceso'}).status_code == 303
    with Session() as db:
        assert db.get(Ticket,2).status == 'en_proceso'


def test_provider_admin_can_reset_and_revoke_technician(sandbox):
    from fastapi.testclient import TestClient
    from app.main import app
    client, Session = sandbox
    signin(client, 'root@hospitalops.example')
    add_technician(client)
    with Session() as db:
        tech_id = db.scalar(select(User).where(User.email == 'pedro@servicios.example')).id
    with TestClient(app, follow_redirects=False) as tech_client:
        sign_in_new(tech_client, 'pedro@servicios.example','PasswordUnicaComercial2026!')
        assert tech_client.get('/mis-repuestos').status_code == 200
        assert form_post(client, f'/equipo/{tech_id}/reset-password', {'password':'ContrasenaRestablecida2026!'}).status_code == 303
        assert tech_client.get('/mis-repuestos').status_code == 303
        tech_client.cookies.clear()
        assert sign_in_new(tech_client, 'pedro@servicios.example','PasswordUnicaComercial2026!').headers['location'] == '/login'
        assert sign_in_new(tech_client, 'pedro@servicios.example','ContrasenaRestablecida2026!').headers['location'] == '/dashboard'
