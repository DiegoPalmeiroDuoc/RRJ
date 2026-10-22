import io
from openpyxl import load_workbook
from sqlalchemy import select
from app.models import AuditLog, User
from app.security import hash_password
from test_app import sandbox, signin, form_post


def user_id(Session, email):
    with Session() as db:
        return db.scalar(select(User).where(User.email == email)).id


def test_logout_revokes_copied_session_cookie(sandbox):
    client, _ = sandbox
    signin(client, 'admin@norte.example')
    copied = dict(client.cookies)
    assert form_post(client, '/logout', {}).status_code == 303
    client.cookies.clear()
    for name, value in copied.items():
        client.cookies.set(name, value)
    response = client.get('/dashboard')
    assert response.status_code == 303 and response.headers['location'] == '/login'


def test_logout_then_login_again_works(sandbox):
    client, _ = sandbox
    signin(client, 'admin@norte.example')
    form_post(client, '/logout', {})
    signin(client, 'admin@norte.example')
    assert client.get('/dashboard').status_code == 200


def test_editing_user_keeps_approval_decision(sandbox):
    client, Session = sandbox
    signin(client, 'admin@norte.example')
    pending = user_id(Session, 'pendiente@norte.example')
    form = {'name': 'Antonia Pendiente', 'email': 'pendiente@norte.example', 'role': 'solicitante', 'job_title': 'TENS'}
    assert form_post(client, f'/users/{pending}/edit', form).status_code == 303
    with Session() as db:
        assert db.get(User, pending).approval == 'pending'
    form_post(client, f'/approvals/{pending}/reject', {})
    assert form_post(client, f'/users/{pending}/edit', form).status_code == 303
    with Session() as db:
        assert db.get(User, pending).approval == 'rejected'


def test_hospital_admin_cannot_manage_other_admin(sandbox):
    client, Session = sandbox
    with Session() as db:
        norte = db.get(User, user_id(Session, 'admin@norte.example')).hospital_id
        db.add(User(name='Segundo Admin', email='admin2@norte.example', role='admin_cliente', hospital_id=norte,
                    approval='approved', password_hash=hash_password('ClaveSegundoAdmin2026')))
        db.commit()
    other = user_id(Session, 'admin2@norte.example')
    signin(client, 'admin@norte.example')
    assert client.get(f'/users/{other}/edit').status_code == 403
    takeover = {'name': 'Segundo Admin', 'email': 'admin2@norte.example', 'role': 'solicitante', 'password': 'ClaveTomada2026!'}
    assert form_post(client, f'/users/{other}/edit', takeover).status_code == 403
    assert form_post(client, f'/users/{other}/toggle', {}).status_code == 403
    with Session() as db:
        target = db.get(User, other)
        assert target.role == 'admin_cliente' and target.active
    # Sigue pudiendo gestionar a usuarios no administradores y su propia ficha
    requester = user_id(Session, 'solicitante@norte.example')
    assert client.get(f'/users/{requester}/edit').status_code == 200
    assert client.get(f"/users/{user_id(Session, 'admin@norte.example')}/edit").status_code == 200
    # El administrador de plataforma sí puede gestionarlo
    client.cookies.clear()
    signin(client, 'root@hospitalops.example')
    assert form_post(client, f'/users/{other}/toggle', {}).status_code == 303


def test_company_audit_is_superadmin_only_and_filters(sandbox):
    client, Session = sandbox
    signin(client, 'root@hospitalops.example')
    assert form_post(client, '/bodega/crear', {'sku': 'AUD-01', 'name': 'Repuesto auditado',
                                               'stock': '3', 'gross_cost': '10', 'sale_price': '20'}).status_code == 303
    page = client.get('/auditoria-empresa')
    assert page.status_code == 200 and 'AUD-01' in page.text
    assert 'Instalación de datos iniciales' not in page.text  # evento de hospital: no aparece en "solo empresa"
    assert 'Instalación de datos iniciales' in client.get('/auditoria-empresa?alcance=todos').text
    assert 'AUD-01' not in client.get('/auditoria-empresa?alcance=1').text
    assert 'AUD-01' not in client.get('/auditoria-empresa?accion=editar').text
    assert client.get('/auditoria-empresa?desde=no-es-fecha&page=-4').status_code == 200
    export = client.get('/auditoria-empresa/export/xlsx')
    assert export.status_code == 200
    rows = list(load_workbook(io.BytesIO(export.content)).active.iter_rows(min_row=2, values_only=True))
    assert any(row[7] == 'AUD-01' and row[1] == 'Empresa' for row in rows)
    with Session() as db:
        assert db.scalar(select(AuditLog).where(AuditLog.entity == 'auditoria', AuditLog.action == 'exportar'))
    for email in ('admin@norte.example', 'coordinador@norte.example'):
        client.cookies.clear()
        signin(client, email)
        assert client.get('/auditoria-empresa').status_code == 403
        assert client.get('/auditoria-empresa/export/xlsx').status_code == 403
